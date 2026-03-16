"""Tests for marksmith.template (requires marksmith[template] extra)."""

import pytest

# Skip the entire module if docxtpl is not installed.
pytest.importorskip("docxtpl")

from docx import Document  # noqa: E402

from marksmith.convert import md_to_docx  # noqa: E402
from marksmith.template import md_to_docx_templated  # noqa: E402

# ── Helpers ───────────────────────────────────────────────────────────────────


def _make_template(tmp_path, paragraphs: list[str]) -> str:
    """Create a minimal .docx template with the given paragraph texts."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / "template.docx"
    doc.save(str(path))
    return str(path)


# ── md_to_docx_templated ──────────────────────────────────────────────────────


def test_template_not_found(tmp_path):
    md_file = tmp_path / "test.md"
    md_file.write_text("# Hello")
    with pytest.raises(FileNotFoundError, match="missing.docx"):
        md_to_docx_templated(str(md_file), str(tmp_path / "out.docx"), "missing.docx")


def test_template_basic_conversion(tmp_path):
    template = _make_template(
        tmp_path,
        [
            "{{ title }}",
            "{{p marksmith_content }}",
        ],
    )
    md_file = tmp_path / "test.md"
    md_file.write_text("---\ntitle: My Doc\n---\n# Hello\n\nSome content.")
    output = tmp_path / "output.docx"
    md_to_docx_templated(str(md_file), str(output), template)
    assert output.exists()
    assert output.stat().st_size > 0


def test_template_all_metadata_placeholders(tmp_path):
    template = _make_template(
        tmp_path,
        [
            "{{ title }}",
            "{{ version }}",
            "{{ author }}",
            "{{ date }}",
            "{{ classification }}",
            "{{p marksmith_content }}",
        ],
    )
    md_file = tmp_path / "test.md"
    md_file.write_text(
        "---\n"
        "title: Full Test\n"
        "version: 2.0\n"
        "author: Paul Cummings\n"
        "date: 2026-03-16\n"
        "classification: Internal\n"
        "---\n\n"
        "# Section\n\nBody text.\n"
    )
    output = tmp_path / "output.docx"
    md_to_docx_templated(str(md_file), str(output), template)
    assert output.exists()


def test_template_no_frontmatter(tmp_path):
    """Template renders without error even if markdown has no front-matter."""
    template = _make_template(
        tmp_path,
        [
            "{{p marksmith_content }}",
        ],
    )
    md_file = tmp_path / "test.md"
    md_file.write_text("# No metadata\n\nJust content.")
    output = tmp_path / "output.docx"
    md_to_docx_templated(str(md_file), str(output), template)
    assert output.exists()


def test_template_creates_parent_dirs(tmp_path):
    template = _make_template(tmp_path, ["{{p marksmith_content }}"])
    md_file = tmp_path / "test.md"
    md_file.write_text("# Hello")
    output = tmp_path / "deep" / "nested" / "output.docx"
    md_to_docx_templated(str(md_file), str(output), template)
    assert output.exists()


def test_template_rich_content(tmp_path):
    """Smoke test: all block elements through a template."""
    template = _make_template(
        tmp_path,
        [
            "{{ title }} - {{ version }}",
            "{{p marksmith_content }}",
            "End of document.",
        ],
    )
    md = (
        "---\ntitle: Rich Doc\nversion: 1.0\n---\n\n"
        "# Heading\n\n"
        "A paragraph with **bold** and *italic*.\n\n"
        "- item one\n- item two\n\n"
        "1. first\n2. second\n\n"
        "> A blockquote.\n\n"
        "```python\nprint('hello')\n```\n\n"
        "| A | B |\n| --- | --- |\n| 1 | 2 |\n\n"
        "---\n"
    )
    md_file = tmp_path / "rich.md"
    md_file.write_text(md)
    output = tmp_path / "rich_output.docx"
    md_to_docx_templated(str(md_file), str(output), template)
    assert output.exists()
    assert output.stat().st_size > 0


# ── md_to_docx dispatch via --template ───────────────────────────────────────


def test_md_to_docx_dispatches_to_template(tmp_path):
    """md_to_docx with template_path routes to the template engine."""
    template = _make_template(
        tmp_path,
        [
            "{{ title }}",
            "{{p marksmith_content }}",
        ],
    )
    md_file = tmp_path / "test.md"
    md_file.write_text("---\ntitle: Dispatch Test\n---\n# Hello")
    output = tmp_path / "output.docx"
    md_to_docx(str(md_file), str(output), template_path=template)
    assert output.exists()


def test_md_to_docx_template_not_found_via_public_api(tmp_path):
    """FileNotFoundError propagates through md_to_docx for missing template."""
    md_file = tmp_path / "test.md"
    md_file.write_text("# Hello")
    with pytest.raises(FileNotFoundError, match="nonexistent_template.docx"):
        md_to_docx(
            str(md_file),
            str(tmp_path / "output.docx"),
            template_path="nonexistent_template.docx",
        )
